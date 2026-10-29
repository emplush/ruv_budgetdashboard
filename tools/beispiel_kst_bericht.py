"""Erzeugt tests/daten/kst-bericht-beispiel.xlsx: KST-Bericht im Aufbau des Konzern-Controllings mit erfundenen Zahlen
(Kostenstelle 413, ohne Namen). Für Tests und zum Ausprobieren des IST-Imports. Aufruf: python3 tools/beispiel_kst_bericht.py"""
import os
import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "DYN Kostenstellenbericht"
ws["G3"] = "1 KST Bericht (vt) 413 PR"
ws["L3"] = "Konzern-Controlling"
ws["L4"] = "Datenstand: 22.10.2026\nAbrechnungsmonat:  09.2026"
ws["A9"] = "Text"
ws["B9"] = "Beispiel (KST141300 Beispielperson)_vt"
for c in "BCD":
    ws[c + "16"] = "Ist"
ws["B18"], ws["C18"], ws["D18"] = "2025 (vt)", "2025 (vt)", "2026 (vt)"
ws["B19"], ws["C19"], ws["D19"] = "Dezember (kum)", "aktMon (kum vt)", "aktMon (kum vt)"
zeilen = [
    ("Personalkosten AAD (000074 [2])_vt", 400000.00, 300000.00, 310000.50),
    ("Personalkosten ID (000006 [2])_vt", 200000.00, 150000.00, 160000.25),
    ("Sachkosten (000145 [2])_vt", None, None, None),
    ("DV-Kosten (000054 [2])_vt", 10000.00, 7000.00, 8000.00),
    ("Raumkosten (000026 [2])_vt", None, None, None),
    ("Werbungs-/Repräsentationskosten (000050 [2])_vt", 50000.00, 40000.00, 30000.10),
    ("Ext. Beratungskosten - ohne DV (001537 [2])_vt", 5000.00, None, 2500.00),
    ("Reise-/Kfz-/Schulungskosten (000045 [2])_vt", 80000.00, 60000.00, 55000.00),
    ("Porto-/Frachtkosten (000025 [2])_vt", None, None, None),
    ("Gen.-/Vers.-/Verb.-Beiträge (001429 [2])_vt", None, None, None),
    ("Abschreibungen (ohne HW/SW) (000048 [2])_vt", 1000.00, 750.00, 700.00),
    ("Personalbeschaffungskosten (001010 [2])_vt", None, None, None),
    ("Büromat./Literatur/Druckstücke (001065 [2])_vt", 100.00, None, 50.00),
    ("Sonstige Sachkosten (000028 [2])_vt", None, None, 120.00),
]
for i, (text, *werte) in enumerate(zeilen):
    r = 20 + i
    ws["A%d" % r] = text
    for c, w in zip("BCD", werte):
        if w is not None:
            ws["%s%d" % (c, r)] = w
for c in "BCD":
    sach = sum(ws["%s%d" % (c, r)].value or 0 for r in range(23, 34))
    ws["%s22" % c] = round(sach, 2)
    ws["%s34" % c] = round(sach + (ws["%s20" % c].value or 0) + (ws["%s21" % c].value or 0), 2)
ws["A34"] = "Primärkosten I (000015 [2])_vt"
ws["A35"] = "MA Gesamt (002900 [950])_vt"
for c, w in zip("BCD", (8, 8, 9)):
    ws["%s35" % c] = w
ziel = os.path.join(os.path.dirname(__file__), "..", "tests", "daten", "kst-bericht-beispiel.xlsx")
wb.save(ziel)
print(ziel)
